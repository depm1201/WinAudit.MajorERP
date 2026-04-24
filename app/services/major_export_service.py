from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationError
from app.core.tenant_session import resolve_tenant_code
from app.repositories.major_repository import MajorRepository
from app.storage.major_storage_manager import MajorStorageManager


@dataclass
class MajorExportSummary:
    batch_id: int
    file_path: str
    file_name: str
    sheets: int


class MajorExportService:
    def __init__(self) -> None:
        self._repo = MajorRepository()
        self._storage = MajorStorageManager()

    def export_batch(self, session: Session, batch_id: int, tenant_id: int) -> MajorExportSummary:
        batch = self._repo.get_import_batch(session, batch_id)
        if not batch:
            raise ValidationError("Batch no encontrado.")

        client = self._repo.get_client_by_id(session, batch.client_id)
        if not client:
            raise ValidationError("Cliente no encontrado.")

        tenant_code = resolve_tenant_code(tenant_id)
        fiscal_period = batch.fiscal_period or "UNKNOWN"
        export_dir = self._storage.build_export_path(tenant_code, client.ruc, fiscal_period)
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        file_name = f"major_export_batch_{batch.id}_{timestamp}.xlsx"
        file_path = export_dir / file_name

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Resumen"
        self._build_summary_sheet(summary_sheet, batch, client.name, client.ruc)

        entry_sheet = workbook.create_sheet("Entradas")
        self._build_entries_sheet(entry_sheet, batch, session)

        validation_sheet = workbook.create_sheet("Validaciones")
        self._build_validation_sheet(validation_sheet, batch.id, session)

        reconciliation_sheet = workbook.create_sheet("Conciliacion")
        self._build_reconciliation_sheet(reconciliation_sheet, batch.id, session)

        mappings_sheet = workbook.create_sheet("Mapeos")
        self._build_mappings_sheet(mappings_sheet, client.id, session)

        self._freeze_and_style(workbook)
        self._autosize(workbook)
        workbook.save(file_path)

        return MajorExportSummary(
            batch_id=batch.id,
            file_path=str(file_path),
            file_name=file_name,
            sheets=len(workbook.sheetnames),
        )

    def _build_summary_sheet(self, sheet, batch, client_name: str, client_ruc: str) -> None:
        rows = [
            ["Batch ID", batch.id],
            ["Cliente", client_name],
            ["RUC", client_ruc],
            ["Tipo de lote", batch.batch_type],
            ["Sistema origen", batch.source_system],
            ["Codigo reporte", batch.report_code],
            ["Nombre reporte", batch.report_name],
            ["Periodo fiscal", batch.fiscal_period],
            ["Estado validacion", batch.validation_status],
            ["Mensaje validacion", batch.validation_message],
            ["Archivo original", batch.original_file_name],
            ["Ruta almacenada", batch.stored_file_path],
            ["Registros", batch.row_count],
            ["Debe total", float(batch.debit_total or 0)],
            ["Haber total", float(batch.credit_total or 0)],
            ["Saldo total", float(batch.balance_total or 0)],
        ]
        sheet["A1"] = "Resumen del lote del mayor"
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A3"] = "Campo"
        sheet["B3"] = "Valor"
        self._style_header(sheet, 3)
        for idx, (label, value) in enumerate(rows, start=4):
            sheet[f"A{idx}"] = label
            sheet[f"B{idx}"] = value
        sheet.freeze_panes = "A3"

    def _build_entries_sheet(self, sheet, batch, session: Session) -> None:
        if batch.report_code == "ACCOUNT_CATALOG":
            headers = [
                "ID",
                "Fila origen",
                "Codigo cuenta",
                "Nombre cuenta",
                "Padre",
                "Nivel",
                "Ruta",
                "Es grupo",
                "Hoja",
                "Archivo",
                "Categoria tributaria",
                "Tipo tributario",
            ]
            rows = self._repo.list_account_rows(session, batch.id)
            data = [
                [
                    row.id,
                    row.source_row,
                    row.account_code,
                    row.account_name,
                    row.parent_account_code,
                    row.account_level,
                    row.account_path,
                    row.is_group,
                    row.source_sheet,
                    row.source_file,
                    row.tax_category,
                    row.tax_type,
                ]
                for row in rows
            ]
        else:
            headers = [
                "ID",
                "Fila origen",
                "Tipo fila",
                "Codigo cuenta",
                "Nombre cuenta",
                "Fecha transaccion",
                "Descripcion",
                "Debe",
                "Haber",
                "Saldo",
                "Documento relacionado",
                "Centro de costo",
                "Proyecto",
                "Persona",
                "Persona cruza",
                "Categoria tributaria",
                "Asiento",
                "RUC",
                "Archivo",
                "Hoja",
            ]
            rows = self._repo.list_ledger_rows(session, batch.id)
            data = [
                [
                    row.id,
                    row.source_row,
                    row.row_type,
                    row.account_code,
                    row.account_name,
                    row.transaction_date,
                    row.description,
                    float(row.debit or 0) if row.debit is not None else None,
                    float(row.credit or 0) if row.credit is not None else None,
                    float(row.balance or 0) if row.balance is not None else None,
                    row.related_document,
                    row.cost_center,
                    row.project,
                    row.person_name,
                    row.cross_ref_person,
                    row.tax_category,
                    row.journal_entry,
                    row.tax_id,
                    row.source_file,
                    row.source_sheet,
                ]
                for row in rows
            ]

        self._write_table(sheet, headers, data)

    def _build_validation_sheet(self, sheet, batch_id: int, session: Session) -> None:
        headers = [
            "ID",
            "Codigo validacion",
            "Nombre validacion",
            "Scope",
            "Cuenta origen",
            "Fila origen",
            "Esperado",
            "Actual",
            "Diferencia",
            "Estado",
            "Detalle",
        ]
        rows = self._repo.list_validation_rows(session, batch_id, scope="MAYOR")
        data = [
            [
                row.id,
                row.validation_code,
                row.validation_name,
                row.validation_scope,
                row.source_account_code,
                row.source_row,
                float(row.expected_value or 0) if row.expected_value is not None else None,
                float(row.actual_value or 0) if row.actual_value is not None else None,
                float(row.difference_value or 0) if row.difference_value is not None else None,
                row.result_status,
                json.dumps(row.detail_json, ensure_ascii=False) if row.detail_json is not None else None,
            ]
            for row in rows
        ]
        self._write_table(sheet, headers, data)

    def _build_reconciliation_sheet(self, sheet, batch_id: int, session: Session) -> None:
        headers = [
            "ID",
            "Codigo validacion",
            "Nombre validacion",
            "Scope",
            "Cuenta origen",
            "Fila origen",
            "Esperado",
            "Actual",
            "Diferencia",
            "Estado",
            "Detalle",
        ]
        rows = self._repo.list_validation_rows(session, batch_id, scope="RECONCILIATION")
        data = [
            [
                row.id,
                row.validation_code,
                row.validation_name,
                row.validation_scope,
                row.source_account_code,
                row.source_row,
                float(row.expected_value or 0) if row.expected_value is not None else None,
                float(row.actual_value or 0) if row.actual_value is not None else None,
                float(row.difference_value or 0) if row.difference_value is not None else None,
                row.result_status,
                json.dumps(row.detail_json, ensure_ascii=False) if row.detail_json is not None else None,
            ]
            for row in rows
        ]
        self._write_table(sheet, headers, data)

    def _build_mappings_sheet(self, sheet, client_id: int, session: Session) -> None:
        headers = ["ID", "Retencion %", "Cuenta destino", "Tipo tributario", "Estado"]
        rows = self._repo.list_account_mappings(session, client_id)
        data = [
            [
                row.id,
                float(row.retention_percentage) if row.retention_percentage is not None else None,
                row.target_account_code,
                row.tax_type,
                row.status,
            ]
            for row in rows
        ]
        self._write_table(sheet, headers, data)

    def _write_table(self, sheet, headers: list[str], data: list[list[object]]) -> None:
        sheet.append(headers)
        for row in data:
            sheet.append(row)
        self._style_header(sheet, 1)
        sheet.freeze_panes = "A2"
        if not data:
            sheet.append(["Sin registros"] + [None] * (len(headers) - 1))

    def _style_header(self, sheet, row_number: int) -> None:
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)
        border = Border(bottom=Side(style="thin", color="D9E1F2"))
        for row in sheet.iter_rows(min_row=row_number, max_row=row_number):
            for cell in row:
                cell.fill = fill
                cell.font = font
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

    def _freeze_and_style(self, workbook: Workbook) -> None:
        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False

    def _autosize(self, workbook: Workbook) -> None:
        for sheet in workbook.worksheets:
            widths: dict[int, int] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = str(cell.value)
                    widths[cell.column] = max(widths.get(cell.column, 0), len(value))
            for column, width in widths.items():
                sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 45)
