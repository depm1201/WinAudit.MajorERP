from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


HEADER_ALIASES = {
    "fecha": "transaction_date",
    "codigo": "account_code",
    "código": "account_code",
    "cuenta": "account_name",
    "centro de costo": "cost_center",
    "proyecto": "project",
    "asiento": "journal_entry",
    "documento": "related_document",
    "identificacion": "tax_id",
    "identificación": "tax_id",
    "persona": "person_name",
    "persona cruce cuenta": "cross_ref_person",
    "descripcion": "description",
    "descripción": "description",
    "debe": "debit",
    "haber": "credit",
    "saldo": "balance",
}


@dataclass
class ParsedAccountRow:
    account_code: str
    account_name: str
    parent_account_code: str | None
    account_level: int | None
    account_path: str | None
    sort_order: int | None
    source_sheet: str | None
    source_row: int | None
    is_group: bool = False
    source_file: str | None = None
    tax_category: str | None = None
    tax_type: str | None = None


@dataclass
class ParsedLedgerRow:
    row_type: str
    account_code: str
    account_name: str | None
    transaction_date: date | None
    description: str | None
    debit: Decimal | None
    credit: Decimal | None
    balance: Decimal | None
    related_document: str | None
    cost_center: str | None
    project: str | None
    person_name: str | None
    cross_ref_person: str | None
    tax_category: str | None
    journal_entry: str | None
    tax_id: str | None
    source_file: str | None
    source_sheet: str | None
    source_row: int | None


@dataclass
class ParsedMajorWorkbook:
    report_type: str
    sheet_name: str
    company_name: str | None
    report_name: str | None
    fiscal_period: str | None
    account_rows: list[ParsedAccountRow] = field(default_factory=list)
    ledger_rows: list[ParsedLedgerRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return len(self.account_rows) + len(self.ledger_rows)

    def ledger_totals(self) -> tuple[Decimal, Decimal, Decimal]:
        debit = sum((row.debit or Decimal("0")) for row in self.ledger_rows)
        credit = sum((row.credit or Decimal("0")) for row in self.ledger_rows)
        balance = self.ledger_rows[-1].balance if self.ledger_rows else Decimal("0")
        return debit, credit, balance or Decimal("0")


class MajorExcelParser:
    def parse(self, file_path: str | Path, fiscal_period_hint: str | None = None) -> ParsedMajorWorkbook:
        wb = load_workbook(file_path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        sheet_name = ws.title

        company_name = self._clean_text(ws["A1"].value)
        report_name = self._clean_text(ws["A2"].value) or sheet_name
        text_block = " ".join(
            self._clean_text(ws[f"A{row}"].value) or ""
            for row in (1, 2, 3)
        )
        fiscal_period = fiscal_period_hint or self._extract_period(text_block)

        if sheet_name.strip().lower() == "plancuentas":
            account_rows = self._parse_account_catalog(ws, source_file=str(Path(file_path).name))
            self._mark_group_accounts(account_rows)
            return ParsedMajorWorkbook(
                report_type="ACCOUNT_CATALOG",
                sheet_name=sheet_name,
                company_name=company_name,
                report_name=report_name,
                fiscal_period=fiscal_period,
                account_rows=account_rows,
            )

        header_row = self._find_header_row(ws)
        if header_row is None:
            raise ValueError("No se encontro la fila de encabezados obligatorios en la hoja del mayor.")

        ledger_rows = self._parse_ledger_rows(ws, header_row, source_file=str(Path(file_path).name))
        return ParsedMajorWorkbook(
            report_type="LEDGER",
            sheet_name=sheet_name,
            company_name=company_name,
            report_name=report_name,
            fiscal_period=fiscal_period,
            ledger_rows=ledger_rows,
        )

    def _parse_account_catalog(self, ws, *, source_file: str) -> list[ParsedAccountRow]:
        rows: list[ParsedAccountRow] = []
        for row_idx in range(4, ws.max_row + 1):
            code = self._clean_text(ws.cell(row_idx, 1).value)
            name = self._clean_text(ws.cell(row_idx, 2).value)
            if not code and not name:
                continue
            if not code or not name:
                continue
            rows.append(
                ParsedAccountRow(
                    account_code=code,
                    account_name=name,
                    parent_account_code=self._parent_account_code(code),
                    account_level=len(code.split(".")),
                    account_path=code,
                    sort_order=row_idx,
                    source_sheet=ws.title,
                    source_row=row_idx,
                    source_file=source_file,
                    tax_category=self._clean_text(ws.cell(row_idx, 3).value),
                )
            )
        return rows

    def _mark_group_accounts(self, rows: list[ParsedAccountRow]) -> None:
        codes = {row.account_code for row in rows}
        child_parents = {row.parent_account_code for row in rows if row.parent_account_code}
        for row in rows:
            row.is_group = row.account_code in child_parents or any(
                other != row.account_code and other.startswith(f"{row.account_code}.")
                for other in codes
            )

    def _parse_ledger_rows(self, ws, header_row: int, *, source_file: str) -> list[ParsedLedgerRow]:
        headers = self._build_header_map(ws, header_row)
        rows: list[ParsedLedgerRow] = []
        for row_idx in range(header_row + 1, ws.max_row + 1):
            raw = {alias: self._clean_cell(ws.cell(row_idx, col).value) for alias, col in headers.items()}
            if not any(raw.values()):
                continue

            account_code = self._clean_text(raw.get("account_code"))
            if not account_code:
                continue

            debit = self._parse_decimal(raw.get("debit"))
            credit = self._parse_decimal(raw.get("credit"))
            balance = self._parse_decimal(raw.get("balance"))
            transaction_date = self._parse_date(raw.get("transaction_date"))

            row_type = self._classify_row_type(transaction_date, debit, credit, balance, raw)

            rows.append(
                ParsedLedgerRow(
                    row_type=row_type,
                    account_code=account_code,
                    account_name=self._clean_text(raw.get("account_name")),
                    transaction_date=transaction_date,
                    description=self._clean_text(raw.get("description")),
                    debit=debit,
                    credit=credit,
                    balance=balance,
                    related_document=self._clean_text(raw.get("related_document")),
                    cost_center=self._clean_text(raw.get("cost_center")),
                    project=self._clean_text(raw.get("project")),
                    person_name=self._clean_text(raw.get("person_name")),
                    cross_ref_person=self._clean_text(raw.get("cross_ref_person")),
                    tax_category=None,
                    journal_entry=self._clean_text(raw.get("journal_entry")),
                    tax_id=self._clean_text(raw.get("tax_id")),
                    source_file=source_file,
                    source_sheet=ws.title,
                    source_row=row_idx,
                )
            )
        return rows

    def _build_header_map(self, ws, header_row: int) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            header = self._normalize_header(ws.cell(header_row, col_idx).value)
            if header in HEADER_ALIASES:
                mapping[HEADER_ALIASES[header]] = col_idx
        required = {"transaction_date", "account_code", "account_name", "debit", "credit", "balance"}
        if not required.issubset(mapping.keys()):
            missing = ", ".join(sorted(required - mapping.keys()))
            raise ValueError(f"Faltan encabezados obligatorios en la hoja: {missing}")
        return mapping

    def _find_header_row(self, ws) -> int | None:
        for row_idx in range(1, min(ws.max_row, 20) + 1):
            headers = {
                self._normalize_header(ws.cell(row_idx, col).value)
                for col in range(1, min(ws.max_column, 16) + 1)
            }
            if {"fecha", "codigo", "cuenta", "debe", "haber", "saldo"}.issubset(headers):
                return row_idx
        return None

    def _classify_row_type(
        self,
        transaction_date: date | None,
        debit: Decimal | None,
        credit: Decimal | None,
        balance: Decimal | None,
        raw: dict[str, str | None],
    ) -> str:
        description = self._clean_text(raw.get("description")) or ""
        if transaction_date is None and balance is not None and debit is None and credit is None:
            return "OPENING"
        if "ajuste" in description.lower():
            return "ADJUSTMENT"
        return "MOVEMENT"

    def _extract_period(self, text: str) -> str | None:
        if not text:
            return None

        date_match = re.search(r"(\d{2})[/-](\d{2})[/-](\d{4})", text)
        if date_match:
            year = int(date_match.group(3))
            month = int(date_match.group(2))
            return f"{year:04d}-{month:02d}"

        period_match = re.search(r"(\d{4})[/-](\d{1,2})", text)
        if period_match:
            return f"{int(period_match.group(1)):04d}-{int(period_match.group(2)):02d}"

        year_match = re.search(r"(20\d{2})", text)
        if year_match:
            return f"{year_match.group(1)}-01"

        return None

    def _parent_account_code(self, code: str) -> str | None:
        parts = [part for part in code.split(".") if part]
        if len(parts) <= 1:
            return None
        return ".".join(parts[:-1])

    def _parse_decimal(self, value) -> Decimal | None:
        if value is None:
            return None
        if isinstance(value, Decimal):
            return value
        if isinstance(value, (int, float)):
            return Decimal(str(value))
        text = self._clean_text(value)
        if not text:
            return None
        text = text.replace("$", "").replace(" ", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", ".")
        try:
            return Decimal(text)
        except (InvalidOperation, ValueError):
            return None

    def _parse_date(self, value) -> date | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = self._clean_text(value)
        if not text:
            return None
        date_match = re.search(r"(\d{2})[/-](\d{2})[/-](\d{4})", text)
        if date_match:
            day, month, year = map(int, date_match.groups())
            try:
                return date(year, month, day)
            except ValueError:
                return None
        return None

    def _normalize_header(self, value) -> str:
        text = self._clean_text(value) or ""
        normalized = unicodedata.normalize("NFKD", text)
        normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
        return " ".join(normalized.lower().split())

    def _clean_cell(self, value) -> str | None:
        if value is None:
            return None
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return self._clean_text(value)

    def _clean_text(self, value) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

